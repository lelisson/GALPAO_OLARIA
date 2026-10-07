# -*- coding: utf-8 -*-
"""Leitura e gravação da planilha modelo (Cadastro, Lançamentos, Faturamento, Orçado)."""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from config import (
    CAD_FIRST_DATA_ROW,
    CAD_HEADER_ROW,
    FAT_FIRST_DATA_ROW,
    FAT_HEADER_ROW,
    FAT_TOTAL_ROW,
    LANC_COL_CATEGORIA,
    LANC_COL_DESPESA,
    LANC_COL_FIRST_MES,
    LANC_COL_TOTAL,
    LANC_FIRST_DATA_ROW,
    LANC_HEADER_ROW,
    MESES,
    PLANILHA_PADRAO,
)
from orcamento import DespesaCadastro


def _f(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def caminho_planilha(path: Path | None = None) -> Path:
    p = path or PLANILHA_PADRAO
    if not p.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {p}")
    return p


def mtime_planilha(path: Path | None = None) -> float:
    return caminho_planilha(path).stat().st_mtime


def carregar_cadastro(path: Path | None = None) -> tuple[int, list[DespesaCadastro]]:
    wb = load_workbook(caminho_planilha(path), read_only=True, data_only=True)
    ws = wb["Cadastro"]
    ano = 2026
    for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
        for cell in row:
            if isinstance(cell, (int, float)) and 2020 <= cell <= 2035:
                ano = int(cell)
    itens: list[DespesaCadastro] = []
    for row in ws.iter_rows(min_row=CAD_FIRST_DATA_ROW, values_only=True):
        if not row or not row[1]:
            break
        if str(row[1]).strip().upper() == "TOTAL":
            break
        itens.append(
            DespesaCadastro(
                numero=int(row[0]) if row[0] is not None else None,
                nome=str(row[1]).strip(),
                categoria=str(row[2] or "").strip(),
                valor=float(row[3]) if row[3] not in (None, "") else None,
                periodicidade=str(row[4]).strip() if row[4] else None,
                mes_pagamento=int(row[5]) if row[5] not in (None, "") else None,
                intervalo_meses=int(row[6]) if row[6] not in (None, "") else None,
                equivalente_mensal=float(row[7]) if row[7] not in (None, "") else None,
                observacao=str(row[9]).strip() if len(row) > 9 and row[9] else None,
            )
        )
    wb.close()
    return ano, itens


def _ler_matriz_mensal(ws, first_row: int, stop_label: str = "TOTAL") -> tuple[list[str], list[str], list[list[float]]]:
    despesas: list[str] = []
    categorias: list[str] = []
    valores: list[list[float]] = []
    for row in ws.iter_rows(min_row=first_row, values_only=True):
        if not row or not row[0]:
            break
        nome = str(row[0]).strip()
        if nome.upper().startswith(stop_label):
            break
        despesas.append(nome)
        categorias.append(str(row[1] or "").strip())
        meses = [_f(row[i]) for i in range(2, 14)]
        valores.append(meses)
    return despesas, categorias, valores


def carregar_lancamentos(path: Path | None = None) -> tuple[list[str], list[str], list[list[float]]]:
    wb = load_workbook(caminho_planilha(path), read_only=True, data_only=True)
    out = _ler_matriz_mensal(wb["Lançamentos"], LANC_FIRST_DATA_ROW)
    wb.close()
    return out


def _ler_total_do_mes(ws, first_data_row: int) -> tuple[list[float], float, float]:
    """Lê a linha TOTAL DO MÊS (mesmos números que o Excel / aba DASHBOARD)."""
    for row in ws.iter_rows(min_row=first_data_row, values_only=True):
        if not row or not row[0]:
            continue
        if str(row[0]).strip().upper().startswith("TOTAL DO MÊS"):
            meses = [_f(row[i]) for i in range(2, 14)]
            total_ano = _f(row[14]) if len(row) > 14 else sum(meses)
            media = _f(row[15]) if len(row) > 15 else (total_ano / 12 if total_ano else 0.0)
            return meses, total_ano, media
    return [0.0] * 12, 0.0, 0.0


def carregar_faturamento(path: Path | None = None) -> tuple[list[str], list[list[float]]]:
    wb = load_workbook(caminho_planilha(path), read_only=True, data_only=True)
    ws = wb["Faturamento"]
    linhas: list[str] = []
    vals: list[list[float]] = []
    for row in ws.iter_rows(min_row=FAT_FIRST_DATA_ROW, values_only=True):
        if not row or not row[0]:
            break
        nome = str(row[0]).strip()
        if nome.upper().startswith("TOTAL"):
            break
        linhas.append(nome)
        vals.append([_f(row[i]) for i in range(2, 14)])
    wb.close()
    return linhas, vals


def _carregar_cadastro_ws(ws) -> tuple[int, list[DespesaCadastro]]:
    ano = 2026
    for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
        for cell in row:
            if isinstance(cell, (int, float)) and 2020 <= cell <= 2035:
                ano = int(cell)
    itens: list[DespesaCadastro] = []
    for row in ws.iter_rows(min_row=CAD_FIRST_DATA_ROW, values_only=True):
        if not row or not row[1]:
            break
        if str(row[1]).strip().upper() == "TOTAL":
            break
        itens.append(
            DespesaCadastro(
                numero=int(row[0]) if row[0] is not None else None,
                nome=str(row[1]).strip(),
                categoria=str(row[2] or "").strip(),
                valor=float(row[3]) if row[3] not in (None, "") else None,
                periodicidade=str(row[4]).strip() if row[4] else None,
                mes_pagamento=int(row[5]) if row[5] not in (None, "") else None,
                intervalo_meses=int(row[6]) if row[6] not in (None, "") else None,
                equivalente_mensal=float(row[7]) if row[7] not in (None, "") else None,
                observacao=str(row[9]).strip() if len(row) > 9 and row[9] else None,
            )
        )
    return ano, itens


def snapshot_dados(path: Path | None = None) -> dict:
    p = caminho_planilha(path)
    wb = load_workbook(p, read_only=True, data_only=True)
    ano, cadastro = _carregar_cadastro_ws(wb["Cadastro"])
    despesas, categorias, realizado = _ler_matriz_mensal(wb["Lançamentos"], LANC_FIRST_DATA_ROW)
    despesa_mes, despesa_ano, media_despesa = _ler_total_do_mes(wb["Lançamentos"], LANC_FIRST_DATA_ROW)
    receita_mes, receita_ano, _ = _ler_total_do_mes(wb["Faturamento"], FAT_FIRST_DATA_ROW)
    fat_linhas, fat_vals = [], []
    for row in wb["Faturamento"].iter_rows(min_row=FAT_FIRST_DATA_ROW, values_only=True):
        if not row or not row[0]:
            break
        nome = str(row[0]).strip()
        if nome.upper().startswith("TOTAL"):
            break
        fat_linhas.append(nome)
        fat_vals.append([_f(row[i]) for i in range(2, 14)])
    wb.close()

    por_categoria: dict[str, float] = {}
    for cat, linha in zip(categorias, realizado):
        por_categoria[cat] = por_categoria.get(cat, 0.0) + sum(linha)

    return {
        "path": str(p),
        "mtime": p.stat().st_mtime,
        "ano": ano,
        "cadastro": cadastro,
        "despesas": despesas,
        "categorias": categorias,
        "realizado": realizado,
        "despesa_mes": despesa_mes,
        "despesa_ano": despesa_ano,
        "media_despesa_mes": media_despesa,
        "receita_mes": receita_mes,
        "receita_ano": receita_ano,
        "faturamento_linhas": fat_linhas,
        "por_categoria_ano": por_categoria,
        "meses": MESES,
    }


def salvar_valor_lancamento(
    despesa: str,
    mes_idx: int,
    valor: float | None,
    path: Path | None = None,
) -> None:
    """mes_idx 0=Jan .. 11=Dez. valor None deixa célula em branco."""
    p = caminho_planilha(path)
    backup = p.with_suffix(f".bak_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    shutil.copy2(p, backup)

    wb = load_workbook(p)
    ws = wb["Lançamentos"]
    row_found = None
    for r in range(LANC_FIRST_DATA_ROW, ws.max_row + 1):
        cell = ws.cell(r, LANC_COL_DESPESA).value
        if cell and str(cell).strip() == despesa:
            row_found = r
            break
    if row_found is None:
        wb.close()
        raise ValueError(f"Despesa não encontrada na aba Lançamentos: {despesa}")

    col = LANC_COL_FIRST_MES + mes_idx
    ws.cell(row_found, col, valor if valor is not None else None)

    # Total da linha
    total = 0.0
    tem_valor = False
    for c in range(LANC_COL_FIRST_MES, LANC_COL_FIRST_MES + 12):
        v = ws.cell(row_found, c).value
        if v is not None and v != "":
            total += float(v)
            tem_valor = True
    ws.cell(row_found, LANC_COL_TOTAL, total if tem_valor else None)

    _recalc_totais_lancamentos(ws)
    wb.save(p)
    wb.close()


def atualizar_faturamento_meses(
    linha: str,
    valores_mes: list[float],
    path: Path | None = None,
) -> None:
    """Atualiza 12 meses de uma linha de Faturamento (1 backup)."""
    if len(valores_mes) != 12:
        raise ValueError("valores_mes deve ter 12 posições (Jan–Dez).")
    p = caminho_planilha(path)
    shutil.copy2(p, p.with_suffix(f".bak_{datetime.now():%Y%m%d_%H%M%S}.xlsx"))
    wb = load_workbook(p)
    ws = wb["Faturamento"]
    row_found = None
    for r in range(FAT_FIRST_DATA_ROW, ws.max_row + 1):
        cell = ws.cell(r, 1).value
        if cell and str(cell).strip() == linha:
            row_found = r
            break
    if row_found is None:
        wb.close()
        raise ValueError(f"Linha não encontrada em Faturamento: {linha}")
    for mes_idx, valor in enumerate(valores_mes):
        col = 3 + mes_idx
        ws.cell(row_found, col, valor if valor and valor > 0 else None)
    _recalc_totais_faturamento(ws)
    wb.save(p)
    wb.close()


def salvar_valor_faturamento(
    linha: str,
    mes_idx: int,
    valor: float | None,
    path: Path | None = None,
) -> None:
    p = caminho_planilha(path)
    shutil.copy2(p, p.with_suffix(f".bak_{datetime.now():%Y%m%d_%H%M%S}.xlsx"))
    wb = load_workbook(p)
    ws = wb["Faturamento"]
    row_found = None
    for r in range(FAT_FIRST_DATA_ROW, ws.max_row + 1):
        cell = ws.cell(r, 1).value
        if cell and str(cell).strip() == linha:
            row_found = r
            break
    if row_found is None:
        wb.close()
        raise ValueError(f"Linha não encontrada em Faturamento: {linha}")

    col = 3 + mes_idx
    ws.cell(row_found, col, valor if valor is not None else None)
    _recalc_totais_faturamento(ws)
    wb.save(p)
    wb.close()


def _recalc_totais_lancamentos(ws) -> None:
    for r in range(LANC_FIRST_DATA_ROW, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a and str(a).strip().upper().startswith("TOTAL DO MÊS"):
            for c in range(LANC_COL_FIRST_MES, LANC_COL_FIRST_MES + 12):
                s = 0.0
                ok = False
                for rr in range(LANC_FIRST_DATA_ROW, r):
                    v = ws.cell(rr, c).value
                    if v is not None and v != "":
                        s += float(v)
                        ok = True
                ws.cell(r, c, s if ok else None)
            tot_ano = sum(
                float(ws.cell(r, c).value or 0)
                for c in range(LANC_COL_FIRST_MES, LANC_COL_FIRST_MES + 12)
                if ws.cell(r, c).value not in (None, "")
            )
            ws.cell(r, LANC_COL_TOTAL, tot_ano if tot_ano else None)
            break


def _recalc_totais_faturamento(ws) -> None:
    for r in range(FAT_FIRST_DATA_ROW, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a and str(a).strip().upper().startswith("TOTAL DO MÊS"):
            for c in range(3, 15):
                s = 0.0
                ok = False
                for rr in range(FAT_FIRST_DATA_ROW, r):
                    v = ws.cell(rr, c).value
                    if v is not None and v != "":
                        s += float(v)
                        ok = True
                ws.cell(r, c, s if ok else None)
            break


def sincronizar_orcado(wb) -> None:
    """Reescreve aba Orçado a partir do Cadastro (mesma lógica do modelo)."""
    from orcamento import grade_orcada

    if "Orçado" not in wb.sheetnames or "Cadastro" not in wb.sheetnames:
        return
    _, cad = carregar_cadastro_from_wb(wb)
    grade = grade_orcada(cad)
    ws_l = wb["Lançamentos"]
    ws_o = wb["Orçado"]

    # Copia estrutura de despesas da aba Lançamentos
    for r in range(LANC_FIRST_DATA_ROW, ws_l.max_row + 1):
        nome = ws_l.cell(r, LANC_COL_DESPESA).value
        if not nome or str(nome).strip().upper().startswith("TOTAL"):
            break
        nome_s = str(nome).strip()
        cat = ws_l.cell(r, LANC_COL_CATEGORIA).value
        ws_o.cell(r, 1, nome_s)
        ws_o.cell(r, 2, cat)
        chave = nome_s
        if chave not in grade:
            for k in grade:
                if k.split()[0] == nome_s.split()[0]:
                    chave = k
                    break
        meses = grade.get(chave, [0.0] * 12)
        for i, v in enumerate(meses):
            ws_o.cell(r, 3 + i, v if v else None)
        ws_o.cell(r, 15, sum(meses) if any(meses) else None)

    _recalc_totais_orcado(ws_o)


def carregar_cadastro_from_wb(wb) -> tuple[int, list[DespesaCadastro]]:
    ws = wb["Cadastro"]
    ano = 2026
    itens: list[DespesaCadastro] = []
    for row in ws.iter_rows(min_row=CAD_FIRST_DATA_ROW, values_only=True):
        if not row or not row[1]:
            break
        if str(row[1]).strip().upper() == "TOTAL":
            break
        itens.append(
            DespesaCadastro(
                numero=int(row[0]) if row[0] is not None else None,
                nome=str(row[1]).strip(),
                categoria=str(row[2] or "").strip(),
                valor=float(row[3]) if row[3] not in (None, "") else None,
                periodicidade=str(row[4]).strip() if row[4] else None,
                mes_pagamento=int(row[5]) if row[5] not in (None, "") else None,
                intervalo_meses=int(row[6]) if row[6] not in (None, "") else None,
                equivalente_mensal=float(row[7]) if row[7] not in (None, "") else None,
                observacao=str(row[9]).strip() if len(row) > 9 and row[9] else None,
            )
        )
    return ano, itens


def _recalc_totais_orcado(ws) -> None:
    for r in range(LANC_FIRST_DATA_ROW, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a and str(a).strip().upper().startswith("TOTAL DO MÊS"):
            for c in range(3, 15):
                s = 0.0
                ok = False
                for rr in range(LANC_FIRST_DATA_ROW, r):
                    v = ws.cell(rr, c).value
                    if v is not None and v != "":
                        s += float(v)
                        ok = True
                ws.cell(r, c, s if ok else None)
            tot = sum(float(ws.cell(r, c).value or 0) for c in range(3, 15) if ws.cell(r, c).value not in (None, ""))
            ws.cell(r, 15, tot if tot else None)
            break
